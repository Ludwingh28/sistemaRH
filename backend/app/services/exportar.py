"""Exportación de la asistencia a Excel (.xlsx), con los mismos filtros y datos que la pantalla."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "0F2338"
VERDE = PatternFill("solid", fgColor="D3EEDF")
ROJO = PatternFill("solid", fgColor="FBDADA")
CELESTE = PatternFill("solid", fgColor="DBEAFE")
GRIS = PatternFill("solid", fgColor="CBD5E1")
CABECERA = PatternFill("solid", fgColor=AZUL)
SUBTOTAL = PatternFill("solid", fgColor="E2E8F0")
FINO = Side(style="thin", color="E2E8F0")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CENTRO = Alignment(horizontal="center", vertical="center")

RELLENO = {"P": VERDE, "F": ROJO, "C": CELESTE, "N": GRIS}
SIGNIFICADO = {"P": "Con preventa", "F": "Sin preventa", "C": "Capacitación", "N": "Feriado", "-": "Fuera de contrato / aún no"}


def _fmt(d) -> str:
    return f"{d:%d/%m/%Y}"


def asistencia_xlsx(matriz, filtros: str, usuario: str) -> bytes:
    """`matriz` es la MatrizAsistencia de la pantalla. Los domingos no se incluyen (igual que en pantalla)."""
    visibles = [i for i, d in enumerate(matriz.dias) if not d.domingo]
    hay_extra = any(f.extra > 0 for f in matriz.filas)

    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "ASISTENCIA DE VENDEDORES"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = f"Del {_fmt(matriz.desde)} al {_fmt(matriz.hasta)}"
    ws["A3"] = filtros
    ws["A4"] = f"Generado el {datetime.now():%d/%m/%Y %H:%M} por {usuario}"
    ws["A4"].font = Font(italic=True, color="64748B", size=9)

    fila0 = 6
    columnas = ["CÓDIGO", "VENDEDOR", "ZONA", "EQUIPO"] + [matriz.dias[i].etiqueta.upper() for i in visibles]
    extras = ["HÁBILES", "ASISTIDOS", "FALTAS", "% ASISTENCIA"] + (["EXTRA"] if hay_extra else []) + ["OBSERVACIÓN"]
    for c, titulo in enumerate(columnas + extras, start=1):
        celda = ws.cell(row=fila0, column=c, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF", size=9)
        celda.fill = CABECERA
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE
    ws.row_dimensions[fila0].height = 30

    n_dias = len(visibles)
    col_hab = 4 + n_dias + 1
    for r, f in enumerate(matriz.filas, start=fila0 + 1):
        ws.cell(row=r, column=1, value=f.codigo_vendedor).font = Font(size=9, color="475569")
        ws.cell(row=r, column=2, value=f.nombre)
        ws.cell(row=r, column=3, value=f.zona.upper())
        ws.cell(row=r, column=4, value=f.equipo.upper())
        for j, i in enumerate(visibles):
            codigo = f.celdas[i]
            celda = ws.cell(row=r, column=5 + j, value="" if codigo == "-" else codigo)
            celda.alignment = CENTRO
            celda.font = Font(size=9, bold=codigo == "P", color="166534" if codigo == "P" else "B91C1C" if codigo == "F" else "1E3A8A")
            if codigo in RELLENO:
                celda.fill = RELLENO[codigo]
        ws.cell(row=r, column=col_hab, value=f.habiles)
        ws.cell(row=r, column=col_hab + 1, value=f.asistidos)
        ws.cell(row=r, column=col_hab + 2, value=f.faltas)
        pct = ws.cell(row=r, column=col_hab + 3, value=(f.asistidos / f.habiles) if f.habiles else None)
        pct.number_format = "0.0%"
        siguiente = col_hab + 4
        if hay_extra:
            ws.cell(row=r, column=siguiente, value=f.extra or None)
            siguiente += 1
        ws.cell(row=r, column=siguiente, value=f"RETIRADO EL {_fmt(f.retirado)}" if f.retirado else None)
        for c in range(1, siguiente + 1):
            ws.cell(row=r, column=c).border = BORDE
            if c >= col_hab and c < siguiente:
                ws.cell(row=r, column=c).alignment = CENTRO

    # Fila de totales
    ultima = fila0 + len(matriz.filas)
    if matriz.filas:
        t = ultima + 1
        hab = sum(f.habiles for f in matriz.filas)
        asi = sum(f.asistidos for f in matriz.filas)
        ws.cell(row=t, column=2, value=f"TOTAL ({len(matriz.filas)} vendedores)")
        ws.cell(row=t, column=col_hab, value=hab)
        ws.cell(row=t, column=col_hab + 1, value=asi)
        ws.cell(row=t, column=col_hab + 2, value=hab - asi)
        pct = ws.cell(row=t, column=col_hab + 3, value=(asi / hab) if hab else None)
        pct.number_format = "0.0%"
        for c in range(1, col_hab + 5 + (1 if hay_extra else 0)):
            ws.cell(row=t, column=c).fill = SUBTOTAL
            ws.cell(row=t, column=c).font = Font(bold=True, size=10)
            ws.cell(row=t, column=c).border = BORDE
            if c >= col_hab:
                ws.cell(row=t, column=c).alignment = CENTRO

    # Leyenda
    leyenda = ultima + 3
    ws.cell(row=leyenda, column=1, value="LEYENDA").font = Font(bold=True, size=9)
    for k, (codigo, texto) in enumerate(SIGNIFICADO.items()):
        ws.cell(row=leyenda + 1 + k, column=1, value="" if codigo == "-" else codigo).alignment = CENTRO
        if codigo in RELLENO:
            ws.cell(row=leyenda + 1 + k, column=1).fill = RELLENO[codigo]
        ws.cell(row=leyenda + 1 + k, column=2, value=texto).font = Font(size=9)

    # Formato
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 17
    for j in range(n_dias):
        ws.column_dimensions[get_column_letter(5 + j)].width = 6.5
    for c in range(col_hab, col_hab + 4 + (1 if hay_extra else 0)):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.column_dimensions[get_column_letter(col_hab + 4 + (1 if hay_extra else 0))].width = 24
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=5)
    if matriz.filas:
        ws.auto_filter.ref = f"A{fila0}:{get_column_letter(col_hab + 3)}{ultima}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila0}:{fila0}"

    # Hoja 2: formato largo (una fila por vendedor y día), cómoda para tablas dinámicas
    det = wb.create_sheet("Detalle")
    for c, t in enumerate(["CÓDIGO", "VENDEDOR", "ZONA", "EQUIPO", "FECHA", "ESTADO"], start=1):
        celda = det.cell(row=1, column=c, value=t)
        celda.font = Font(bold=True, color="FFFFFF", size=9)
        celda.fill = CABECERA
    fila = 2
    for f in matriz.filas:
        for i in visibles:
            codigo = f.celdas[i]
            if codigo == "-":
                continue
            det.cell(row=fila, column=1, value=f.codigo_vendedor)
            det.cell(row=fila, column=2, value=f.nombre)
            det.cell(row=fila, column=3, value=f.zona.upper())
            det.cell(row=fila, column=4, value=f.equipo.upper())
            det.cell(row=fila, column=5, value=matriz.dias[i].fecha).number_format = "DD/MM/YYYY"
            det.cell(row=fila, column=6, value=SIGNIFICADO[codigo].upper())
            fila += 1
    for col, ancho in zip("ABCDEF", (9, 34, 14, 17, 12, 16)):
        det.column_dimensions[col].width = ancho
    det.freeze_panes = "A2"
    det.auto_filter.ref = f"A1:F{max(fila - 1, 1)}"

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
